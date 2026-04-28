"""Word 文件读取工具"""
from docx import Document
from typing import Optional
import io


def read_word(file_content: bytes) -> str:
    """
    读取 Word 文件内容
    
    Args:
        file_content: Word 文件的二进制内容
        
    Returns:
        提取的文本内容
    """
    try:
        doc_file = io.BytesIO(file_content)
        doc = Document(doc_file)
        
        text_parts = []
        for para in doc.paragraphs:
            if para.text.strip():
                text_parts.append(para.text)
        
        # 读取表格内容
        for table in doc.tables:
            for row in table.rows:
                row_text = []
                for cell in row.cells:
                    if cell.text.strip():
                        row_text.append(cell.text.strip())
                if row_text:
                    text_parts.append(" | ".join(row_text))
        
        return "\n".join(text_parts)
    except Exception as e:
        raise Exception(f"Word 文件读取失败: {str(e)}")


def read_word_from_path(file_path: str) -> str:
    """
    从文件路径读取 Word 文件内容
    
    Args:
        file_path: Word 文件路径
        
    Returns:
        提取的文本内容
    """
    try:
        doc = Document(file_path)
        
        text_parts = []
        for para in doc.paragraphs:
            if para.text.strip():
                text_parts.append(para.text)
        
        # 读取表格内容
        for table in doc.tables:
            for row in table.rows:
                row_text = []
                for cell in row.cells:
                    if cell.text.strip():
                        row_text.append(cell.text.strip())
                if row_text:
                    text_parts.append(" | ".join(row_text))
        
        return "\n".join(text_parts)
    except Exception as e:
        raise Exception(f"Word 文件读取失败: {str(e)}")